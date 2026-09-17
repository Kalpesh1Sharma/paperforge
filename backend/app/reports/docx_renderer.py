"""Editable Word export generated directly from the presentation model."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.reports.exceptions import ReportRenderingError
from app.reports.presentation_models import (
    EvidenceTable,
    PresentationEvidence,
    PresentationModel,
    PresentationSection,
)


class EditableDocxRenderer:
    """Create a normal, fully editable DOCX without external services."""

    def render_presentation(self, presentation: PresentationModel, output_path: Path) -> Path:
        target = Path(output_path)
        try:
            document = Document()
            self._configure(document, presentation)
            self._cover(document, presentation)
            document.add_page_break()
            self._contents(document, presentation)
            document.add_page_break()
            for index, section in enumerate(presentation.sections):
                if index:
                    document.add_page_break()
                self._section(document, section)
            if presentation.bibliography:
                document.add_page_break()
                document.add_heading("Bibliography", level=1)
                for entry in presentation.bibliography:
                    paragraph = document.add_paragraph(entry.citation, style="Bibliography")
                    paragraph.paragraph_format.keep_together = True
            target.parent.mkdir(parents=True, exist_ok=True)
            document.save(target)
        except (OSError, ValueError, TypeError) as exc:
            raise ReportRenderingError("Unable to render the editable DOCX report.") from exc
        return target

    @staticmethod
    def _configure(document: Document, presentation: PresentationModel) -> None:
        section = document.sections[0]
        if presentation.page_size == "A4":
            section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        else:
            section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.85)

        fonts = {
            "paperforge-classic": ("Georgia", "Aptos"),
            "modern-research": ("Aptos Display", "Aptos"),
            "ieee-inspired-technical": ("Times New Roman", "Times New Roman"),
        }
        heading_font, body_font = fonts.get(presentation.template_key, ("Georgia", "Aptos"))
        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = body_font
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.15
        for style_name, size in (("Title", 28), ("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 12)):
            style = styles[style_name]
            style.font.name = heading_font
            style.font.size = Pt(size)
            style.font.bold = style_name != "Title"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.keep_with_next = True
        if "Bibliography" not in styles:
            bibliography = styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
        else:
            bibliography = styles["Bibliography"]
        bibliography.font.name = body_font
        bibliography.font.size = Pt(10.5)
        bibliography.paragraph_format.left_indent = Inches(0.3)
        bibliography.paragraph_format.first_line_indent = Inches(-0.3)
        bibliography.paragraph_format.space_after = Pt(8)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Page ")
        run.font.size = Pt(9)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run._r.addnext(field)

    @staticmethod
    def _cover(document: Document, presentation: PresentationModel) -> None:
        cover = presentation.cover
        document.core_properties.title = cover.title
        document.core_properties.author = cover.author or "PaperForge"
        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title.add_run(cover.title)
        if cover.subtitle:
            subtitle = document.add_paragraph(cover.subtitle)
            subtitle.style = document.styles["Subtitle"]
        document.add_paragraph(cover.publication_type or "Research report")
        table = document.add_table(rows=0, cols=2)
        table.style = "Light Shading Accent 1"
        details = (
            ("Prepared by", cover.author or "PaperForge"),
            ("Organisation", cover.organisation),
            ("University", cover.university),
            ("Department", cover.department),
            ("Research domain", cover.domain),
            ("Prepared from", cover.filename),
            ("Generated", cover.generated_on.isoformat() if cover.generated_on else "Not available"),
            ("Citation style", presentation.citation_style.upper()),
        )
        for label, value in details:
            if not value:
                continue
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, str(value)
            cells[0].vertical_alignment = cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[0].paragraphs[0].runs[0].bold = True

    @staticmethod
    def _contents(document: Document, presentation: PresentationModel) -> None:
        document.add_heading("Table of Contents", level=1)
        for index, entry in enumerate(presentation.table_of_contents.entries, start=1):
            document.add_paragraph(f"{index}. {entry.heading}")
        if presentation.bibliography:
            document.add_paragraph(f"{len(presentation.table_of_contents.entries) + 1}. Bibliography")

    @classmethod
    def _section(cls, document: Document, section: PresentationSection) -> None:
        document.add_heading(section.heading, level=1)
        if section.edited_content is not None:
            for block in section.edited_content.split("\n\n"):
                document.add_paragraph(block.strip())
            return
        for paragraph in section.intro:
            document.add_paragraph(paragraph)
        for group in section.finding_groups:
            if group.heading != "General":
                document.add_heading(group.heading, level=2)
            for finding in group.findings:
                if not finding.summary_includes_title:
                    document.add_heading(finding.title, level=3)
                if finding.summary != finding.title or finding.summary_includes_title:
                    document.add_paragraph(finding.summary)
                cls._provenance(document, finding.evidence)
        for group in section.entity_groups:
            document.add_heading(group.category, level=2)
            for entity in group.entities:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(entity.name).bold = True
                if entity.aliases:
                    paragraph.add_run(f" ({', '.join(entity.aliases)})")
                cls._provenance(document, entity.evidence)
        for concept in section.concepts:
            document.add_heading(concept.concept, level=2)
            document.add_paragraph(concept.definition)
            if concept.why_it_matters:
                document.add_paragraph(f"Why it matters: {concept.why_it_matters}")
            cls._provenance(document, concept.evidence)
        for event in section.timeline:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{event.date}: ").bold = True
            paragraph.add_run(event.description)
            cls._provenance(document, event.evidence)
        for table_model in section.evidence_tables:
            cls._table(document, table_model)
        for appendix in section.appendix_groups:
            document.add_heading(appendix.heading, level=2)
            for finding in appendix.findings:
                document.add_heading(finding.title, level=3)
                document.add_paragraph(finding.summary)
                cls._provenance(document, finding.evidence)
            for concept in appendix.concepts:
                document.add_heading(concept.concept, level=3)
                document.add_paragraph(concept.definition)
                if concept.why_it_matters:
                    document.add_paragraph(f"Why it matters: {concept.why_it_matters}")
                cls._provenance(document, concept.evidence)
            for entity_group in appendix.entities:
                document.add_heading(entity_group.category, level=3)
                for entity in entity_group.entities:
                    paragraph = document.add_paragraph(style="List Bullet")
                    paragraph.add_run(entity.name).bold = True
                    if entity.aliases:
                        paragraph.add_run(f" ({', '.join(entity.aliases)})")
                    cls._provenance(document, entity.evidence)
            for reference in appendix.references:
                document.add_paragraph(reference.reference, style="List Number")
                cls._provenance(document, reference.evidence)
            for table_model in appendix.evidence_tables:
                cls._table(document, table_model)
        for reference in section.references:
            document.add_paragraph(reference.reference, style="List Number")
            cls._provenance(document, reference.evidence)
        if not any((section.edited_content, section.intro, section.finding_groups, section.entity_groups, section.concepts, section.timeline, section.evidence_tables, section.appendix_groups, section.references)):
            document.add_paragraph("No material is available for this section.")

    @staticmethod
    def _table(document: Document, table_model: EvidenceTable) -> None:
        """Render one presentation evidence table as editable Word cells."""
        document.add_heading(table_model.title, level=2)
        table = document.add_table(rows=1, cols=len(table_model.columns))
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, table_model.columns, strict=True):
            cell.text = value
            cell.paragraphs[0].runs[0].bold = True
        for row in table_model.rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row, strict=True):
                cell.text = value

    @staticmethod
    def _provenance(document: Document, evidence: PresentationEvidence) -> None:
        if not evidence.source_labels:
            return
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run("Supported by " + ", ".join(evidence.source_labels))
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(80, 80, 80)
